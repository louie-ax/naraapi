import os
import requests
import re
import fitz  # pymupdf
import pdfplumber
import olefile
import zlib
import zipfile
import xml.etree.ElementTree as ET
import openpyxl # xlsx
import xlrd # xls
import subprocess # ★ LibreOffice 실행용
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from io import BytesIO
from fastapi import FastAPI, Request
from supabase import create_client
from datetime import datetime
from pydantic import BaseModel
from typing import Dict, Any, Optional

# 환경변수
SUPABASE_URL = os.environ.get("SUPABASE_URL")
SUPABASE_KEY = os.environ.get("SUPABASE_KEY")

print("\n========== [Worker v10.0 (Docker + LibreOffice)] ==========")
if SUPABASE_URL: print(f"✅ URL Loaded")
else: print("❌ URL Missing")
print("===========================================================\n")

app = FastAPI()
if SUPABASE_URL and SUPABASE_KEY:
    supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
else:
    supabase = None

class SupabaseWebhook(BaseModel):
    type: str = "INSERT"
    table: str = "bid_notices"
    record: Dict[str, Any]
    schema_name: str = "public"
    old_record: Optional[Dict[str, Any]] = None

# --- [1. LibreOffice 변환 함수 (핵심)] ---
def convert_hwp_to_docx(hwp_bytes):
    filename = "temp.hwp"
    docx_filename = "temp.docx"
    try:
        # 파일 저장
        with open(filename, "wb") as f:
            f.write(hwp_bytes)
        
        # 변환 명령 실행 (HWP -> DOCX)
        # soffice --headless --convert-to docx --outdir . temp.hwp
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "docx", "--outdir", ".", filename],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE
        )
        
        # 변환된 파일 읽기
        if os.path.exists(docx_filename):
            with open(docx_filename, "rb") as f:
                docx_bytes = f.read()
            return docx_bytes
        return None
        
    except Exception as e:
        print(f"⚠️ 변환 실패: {e}")
        return None
    finally:
        # 청소
        if os.path.exists(filename): os.remove(filename)
        if os.path.exists(docx_filename): os.remove(docx_filename)

# --- [기존 파싱 함수들 (v9.0 기능 유지)] ---

# (백업용) 기존 HWP 파서
def get_hwp_text_fallback(file_bytes):
    try:
        f = BytesIO(file_bytes)
        ole = olefile.OleFileIO(f)
        dirs = ole.listdir()
        body_sections = []
        for d in dirs:
            if d[0] == "BodyText": body_sections.append(d)
        body_sections.sort(key=lambda x: int(x[1][7:]))
        raw_text = ""
        for section in body_sections:
            stream = ole.openstream(section)
            data = stream.read()
            try: unpacked = zlib.decompress(data, -15)
            except: unpacked = data
            raw_text += unpacked.decode('utf-16le', errors='ignore')
        
        cleaned = re.sub(r'[^가-힣a-zA-Z0-9\s\.\,\-\(\)\[\]\%\~\:\/]', '', raw_text)
        final = []
        for line in cleaned.split('\n'):
            if len(line.strip()) > 1: final.append(line.strip())
        return "\n".join(final)
    except: return "(HWP 내용 없음)"

def get_hwpx_text(file_bytes):
    try:
        text = ""
        with zipfile.ZipFile(BytesIO(file_bytes)) as zf:
            for name in zf.namelist():
                if name.startswith("Contents/section") and name.endswith(".xml"):
                    root = ET.fromstring(zf.read(name))
                    for t in root.iter():
                        if t.tag.endswith('t') and t.text: text += t.text + "\n"
        return text
    except: return "(HWPX 오류)"

def get_xlsx_text(file_bytes):
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        text = ""
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            text += f"\n### [{sheet}]\n"
            rows = list(ws.iter_rows(values_only=True))
            if not rows: continue
            headers = [str(c).replace('\n',' ').strip() if c else "" for c in rows[0]]
            text += "| " + " | ".join(headers) + " |\n|" + "---|"*len(headers) + "\n"
            for row in rows[1:]:
                cells = [str(c).replace('\n',' ').strip() if c else "" for c in row]
                if any(cells): text += "| " + " | ".join(cells) + " |\n"
            text += "\n"
        return text
    except: return "(XLSX 오류)"

def get_xls_text(file_bytes):
    try:
        wb = xlrd.open_workbook(file_contents=file_bytes)
        text = ""
        for s in wb.sheet_names():
            ws = wb.sheet_by_name(s)
            text += f"\n### [{s}]\n"
            if ws.nrows == 0: continue
            headers = [str(ws.cell_value(0,c)).replace('\n',' ').strip() for c in range(ws.ncols)]
            text += "| " + " | ".join(headers) + " |\n|" + "---|"*len(headers) + "\n"
            for r in range(1, ws.nrows):
                cells = [str(ws.cell_value(r,c)).replace('\n',' ').strip() for c in range(ws.ncols)]
                if any(cells): text += "| " + " | ".join(cells) + " |\n"
            text += "\n"
        return text
    except: return "(XLS 오류)"

# DOCX 파서 (순서 정렬 + 표 마크다운)
def iter_block_items(parent):
    if isinstance(parent, _Document): parent_elm = parent.element.body
    elif isinstance(parent, _Cell): parent_elm = parent._tc
    else: return
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P): yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl): yield Table(child, parent)

def get_docx_text(file_bytes):
    try:
        doc = Document(BytesIO(file_bytes))
        full_text = []
        for block in iter_block_items(doc):
            if isinstance(block, Paragraph):
                if block.text.strip(): full_text.append(block.text)
            elif isinstance(block, Table):
                if not block.rows: continue
                full_text.append("\n")
                rows = []
                for row in block.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    rows.append(cells)
                if not rows: continue
                headers = rows[0]
                full_text.append("| " + " | ".join(headers) + " |")
                full_text.append("| " + " | ".join(["---"]*len(headers)) + " |")
                for r in rows[1:]: full_text.append("| " + " | ".join(r) + " |")
                full_text.append("\n")
        return "\n".join(full_text)
    except Exception as e: return f"(DOCX 오류: {e})"

def get_pdf_text(file_bytes):
    try:
        text_output = []
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                tables = page.find_tables()
                bboxes = [t.bbox for t in tables]
                def not_table(obj):
                    if obj["object_type"] == "char":
                        cx, cy = (obj["x0"]+obj["x1"])/2, (obj["top"]+obj["bottom"])/2
                        for b in bboxes:
                            if b[0]<=cx<=b[2] and b[1]<=cy<=b[3]: return False
                    return True
                
                raw = page.filter(not_table).extract_text(layout=True)
                if raw: text_output.append(raw)
                
                extracted = page.extract_tables()
                if extracted:
                    text_output.append("\n")
                    for tbl in extracted:
                        clean = [[(str(c).replace('\n',' ') if c else "") for c in r] for r in tbl]
                        clean = [r for r in clean if any(r)]
                        if not clean: continue
                        h = clean[0]
                        text_output.append("| " + " | ".join(h) + " |")
                        text_output.append("| " + " | ".join(["---"]*len(h)) + " |")
                        for r in clean[1:]: text_output.append("| " + " | ".join(r) + " |")
                    text_output.append("\n")
        return "\n".join(text_output)
    except: return "(PDF 오류)"

# --- [통합 핸들러 (수정됨)] ---
def extract_text_from_file(file_bytes, ext):
    ext = ext.lower()
    try:
        if 'pdf' in ext: return get_pdf_text(file_bytes)
        elif 'docx' in ext or 'doc' in ext: return get_docx_text(file_bytes)
        elif 'hwpx' in ext: return get_hwpx_text(file_bytes)
        elif 'xlsx' in ext or 'xlsm' in ext: return get_xlsx_text(file_bytes)
        elif 'xls' in ext: return get_xls_text(file_bytes)
        
        # ★ HWP 처리 로직 변경 ★
        elif 'hwp' == ext:
            print("  🔄 [변환] HWP -> DOCX 변환 시도 (LibreOffice)")
            # 1. LibreOffice로 DOCX 변환
            docx_bytes = convert_hwp_to_docx(file_bytes)
            
            if docx_bytes:
                print("  ✨ [성공] DOCX 변환 성공 -> 표 파싱 진행")
                # 2. 변환된 DOCX를 우리가 만든 강력한 DOCX 파서로 처리
                return get_docx_text(docx_bytes)
            else:
                print("  ⚠️ [실패] 변환 실패 -> 기존 방식(텍스트만) 사용")
                return get_hwp_text_fallback(file_bytes)

        else: return f"(지원불가: {ext})"
    except Exception as e: return f"시스템 에러: {str(e)}"

# --- [엔드포인트] ---
@app.post("/parse")
async def parse_notice(payload: SupabaseWebhook):
    if not supabase: return {"status": "Error", "msg": "Env Vars Missing"}
    try:
        record = payload.record
        if not record or record.get('process_status') != 'NEW': return {"msg": "Skipped"}
        
        bid_no = record.get('bidNtceNo')
        print(f"🚀 [시작] {bid_no}")
        supabase.table('bid_notices').update({'process_status': 'PROCESSING'}).eq('bidNtceNo', bid_no).execute()
        
        full_log = ""
        for i in range(1, 11):
            url = record.get(f"ntceSpecDocUrl{i}")
            name = record.get(f"ntceSpecFileNm{i}", f"File_{i}")

            if url and str(url).startswith("http"):
                print(f"  📥 [다운로드] {name}")
                try:
                    resp = requests.get(url, timeout=60)
                    text = extract_text_from_file(resp.content, name.split('.')[-1])
                    supabase.table('bid_attachments').insert({
                        "bidNtceNo": bid_no, "bidNtceOrd": record.get('bidNtceOrd'),
                        "file_name": name, "file_url": url, "file_type": name.split('.')[-1],
                        "extracted_text": text
                    }).execute()
                    full_log += f"  💾 [성공] {name}\n"
                except Exception as e:
                    full_log += f"  ❌ [실패] {name}\n"
                    print(f"에러: {e}")

        supabase.table('bid_notices').update({'process_status': 'DONE', 'parsed_content': full_log, 'updated_at': datetime.now().isoformat()}).eq('bidNtceNo', bid_no).execute()
        return {"status": "Success"}
    except Exception as e: return {"status": "Error", "msg": str(e)}